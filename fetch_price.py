import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import yfinance as yf

# 1. Fetch 5-minute data for NIFTY 50 (^NSEI) for the last 60 days
df = yf.download(
    tickers="^NSEI", period="60d", interval="5m", progress=False
)

if isinstance(df.columns, pd.MultiIndex):
    df.columns = df.columns.get_level_values(0)

# Convert timezone to IST (Asia/Kolkata)
if df.index.tz is None:
    df.index = df.index.tz_localize("UTC").tz_convert("Asia/Kolkata")
else:
    df.index = df.index.tz_convert("Asia/Kolkata")

# 2. Process opening session slice (09:15 to 09:45 AM IST)
records = []
for date, day_df in df.groupby(df.index.date):
    session_slice = day_df.between_time("09:15", "09:45")

    if not session_slice.empty:
        open_price = float(session_slice.iloc[0]["Open"])
        high_price = float(session_slice["High"].max())
        low_price = float(session_slice["Low"].min())
        close_price = float(session_slice.iloc[-1]["Close"])

        session_range = high_price - low_price
        change_pct = ((close_price - open_price) / open_price) * 100

        records.append({
            "Date": str(date),
            "Open (09:15)": round(open_price, 2),
            "High (09:15-09:45)": round(high_price, 2),
            "Low (09:15-09:45)": round(low_price, 2),
            "Close (09:45)": round(close_price, 2),
            "Range (Pts)": round(session_range, 2),
            "Change (%)": round(change_pct, 2),
        })

# 3. Build Excel Workbook with openpyxl styling
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "NIFTY_Opening_Session"
ws.views.sheetView[0].showGridLines = True

# Title banner
ws.merge_cells("A1:G1")
ws["A1"] = "NIFTY 50 — Opening Session Analysis (09:15 - 09:45 IST)"
ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill(
    start_color="1F4E78", end_color="1F4E78", fill_type="solid"
)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 35

# Table Headers
headers = [
    "Date",
    "Open (09:15)",
    "High (09:15-09:45)",
    "Low (09:15-09:45)",
    "Close (09:45)",
    "Range (Pts)",
    "Change (%)",
]
header_row = 4

for col_idx, text in enumerate(headers, 1):
    cell = ws.cell(row=header_row, column=col_idx, value=text)
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(
        start_color="2F5597", end_color="2F5597", fill_type="solid"
    )
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

ws.row_dimensions[header_row].height = 25

# Populate Data Rows
start_row = 5
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

for r_idx, rec in enumerate(records, start=start_row):
    ws.row_dimensions[r_idx].height = 20
    row_vals = list(rec.values())

    for c_idx, val in enumerate(row_vals, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.font = Font(name="Arial", size=10)
        cell.border = thin_border

        if c_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c_idx in range(2, 7):
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif c_idx == 7:
            cell.value = val / 100.0
            cell.number_format = "+0.00%;-0.00%;0.00%"
            cell.alignment = Alignment(horizontal="right", vertical="center")

# Save workbook
file_name = "nifty_opening_session_60d.xlsx"
wb.save(file_name)
print(f"File saved successfully as {file_name}")